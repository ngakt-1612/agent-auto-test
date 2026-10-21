#!/usr/bin/env python3
"""Sinh file TC AI Chat (Genius 2.9.0, spec pageId 323393416 §3–11) cùng khuôn 13 cột với
TC_Genius_v2.9.0_UpdateUI_Library.xlsx. Usage: gen_ai_chat_testcases.py <out.xlsx> <spec-cases.json>"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

OUT, JSON_OUT = sys.argv[1], sys.argv[2]
BASE = "1. App Genius 2.9.0 đã cài, mở app tới Home (301)\n2. RC: enable_ai_chat_feature = true"
SUB = BASE + "\n3. User đã mua gói, còn ≥10 lượt"
NOSUB = BASE + "\n3. User chưa mua gói (gói hết hạn), ví topup = 0"
EMPTY = SUB + "\n4. Cuộc trò chuyện đang trống (đã bấm New conversation)"
HASRES = SUB + "\n4. Cuộc trò chuyện có ≥1 kết quả ảnh do Genius tạo"
TWORES = SUB + "\n4. Cuộc trò chuyện có ≥2 kết quả ảnh do Genius tạo"
VIBES = BASE + "\n3. RC: enable_home_vibes_feature = true\n4. Đã cấp quyền Photos (Full Access)"
A, S = "Auto", "Auto + setup"

# (group, test_description, sub_scenario, pre, data, action, expected, priority, automation, note)
G = {}
def tc(group, desc, sub, pre, data, action, exp, prio, auto=A, note=""):
    G.setdefault(group, []).append((desc, sub, pre, data, action, exp, prio, auto, note))

# ---- CH01 Lối vào Chat & trạng thái trống (§3) --------------------------------------------
g = "AI Chat — Lối vào & trạng thái trống (901)"
tc(g, "UI trạng thái trống", "Bố cục màn Chat trống", EMPTY, "N/A", "1. Chạm tab Chat\n2. Quan sát màn",
   "1. Đầu màn: tiêu đề 'Chat' ở giữa, nút ba chấm góc phải\n2. Giữa màn: câu chào \"What's on your mind?\"\n3. Ngay dưới: 3 chip xếp dọc, canh trái: 'Make a birthday card from my photo' · 'Turn my photo into a film' · 'Try a black-and-white look'\n4. Đáy: thanh soạn gồm nút dấu cộng, chữ mờ \"What's on your mind?\", nút gửi tròn mũi tên lên đang mờ\n5. Không có danh sách tin nhắn", "high", A, "Assert label tiêu đề, câu chào, 3 chip, nút gửi isEnabled=false")
tc(g, "UI trạng thái trống", "Chip gợi ý là chữ, không có biểu tượng", EMPTY, "N/A", "1. Mở màn Chat trống\n2. Quan sát 3 chip",
   "1. Mỗi chip chỉ có chữ, không có icon\n2. Mỗi chip nằm trên một dòng riêng", "low", A, "Đọc cây accessibility: chip không có image con")
tc(g, "Vào Chat", "Bấm tab Chat trên thanh tab", BASE + "\n3. Đang ở tab Home", "N/A", "1. Chạm 'Chat' trên thanh tab",
   "1. Màn Chat (901) mở trong tab Chat\n2. Mục Chat trên thanh tab sáng (tím)\n3. Thanh tab vẫn hiện", "critical", A, "isSelected + pixel tím tab Chat")
tc(g, "Vào Chat", "Bấm tab Chat lần nữa khi đang ở Chat", SUB + "\n4. Cuộc trò chuyện có nhiều tin nhắn, đang cuộn lên đầu", "N/A", "1. Cuộn danh sách lên đầu\n2. Chạm 'Chat' trên thanh tab lần nữa",
   "1. Danh sách cuộn xuống tin nhắn cuối trong ~0,3 giây", "low", S, "Cần cuộc trò chuyện ≥6 tin; so vị trí tin cuối trước/sau")
tc(g, "Vào Chat", "Start creating ở Library trống mở Chat", BASE + "\n3. Library trống (chưa từng tạo gì)", "N/A", "1. Mở tab Library\n2. Bấm 'Start creating'",
   "1. Thanh tab chuyển sang Chat (mục Chat sáng)\n2. Màn Chat (901) hiện", "high", S, "Đã chạy ở TC-LIB01-003 suite Library — chạy lại khi Library trống")
tc(g, "Logic chip gợi ý", "Bấm chip điền vào ô nhập, chưa gửi", EMPTY, "Chip: 'Turn my photo into a film'", "1. Bấm chip 'Turn my photo into a film'",
   "1. Chữ của chip nằm trong ô nhập\n2. Bàn phím mở\n3. Chưa có tin nhắn nào được gửi, số lượt không đổi", "high", A, "Đọc value ô nhập; đếm bong bóng = 0; huy hiệu lượt trước/sau")
tc(g, "Logic lịch sử", "Mở lại app thấy cuộc trò chuyện cũ", HASRES, "N/A", "1. Kill app\n2. Mở lại app\n3. Chạm tab Chat",
   "1. Thấy lại toàn bộ tin nhắn và kết quả cũ\n2. Danh sách đứng ở tin nhắn cuối\n3. Không hiện câu chào", "high", A, "So số tin + chữ tin cuối trước/sau relaunch")
tc(g, "Logic cuộc trò chuyện mới", "New conversation + xác nhận", HASRES, "N/A", "1. Bấm nút ba chấm góc trên\n2. Chọn 'New conversation'\n3. Xác nhận ở hộp thoại",
   "1. Hộp thoại hỏi lại: 'Start a new conversation? Your images and videos stay in Library.'\n2. Danh sách tin nhắn biến mất, màn về trạng thái trống (câu chào + chip)\n3. Ảnh/video đã tạo vẫn còn trong tab Library", "high", A, "[⚠️ AMBIGUOUS C1] spec ghi 'Cuộc trò chuyện mới', app tiếng Anh 'New conversation'")
tc(g, "Logic cuộc trò chuyện mới", "New conversation + Huỷ", HASRES, "N/A", "1. Bấm nút ba chấm\n2. Chọn 'New conversation'\n3. Bấm Huỷ ở hộp thoại",
   "1. Hộp thoại đóng\n2. Cuộc trò chuyện giữ nguyên, đủ tin nhắn", "medium", A, "")
tc(g, "Edge mạng", "Mở Chat khi mất mạng", HASRES + "\n5. Bật Airplane Mode + tắt Wi-Fi", "N/A", "1. Mở tab Chat\n2. Gõ 'hello'",
   "1. Màn Chat mở, thấy cuộc trò chuyện cũ đã lưu trên máy\n2. Gõ được chữ vào ô nhập", "high", S, "Mất mạng qua app Settings (SettingsNetworkDriver)")
tc(g, "Edge cuộn", "Kéo quá mép trên", SUB + "\n4. Cuộc trò chuyện có ≥6 tin", "N/A", "1. Vuốt nhanh lên đầu danh sách\n2. Kéo quá mép trên rồi thả",
   "1. Danh sách dừng ở tin nhắn đầu tiên\n2. Thả tay thì trở về sát mép trên\n3. Không có khoảng trống sau tin nhắn cuối", "low", S, "So frame tin đầu với mép trên sau khi thả")

# ---- CH02 Thanh soạn & gửi (§4) -------------------------------------------------------------
g = "AI Chat — Thanh soạn & gửi yêu cầu (901)"
tc(g, "Validate nút gửi", "Ô nhập trống → nút gửi tắt", EMPTY, "Ô nhập: ''", "1. Quan sát nút gửi\n2. Bấm nút gửi",
   "1. Nút gửi mờ, không bấm được\n2. Không có tin nhắn nào được gửi", "high", A, "isEnabled=false; đếm bong bóng")
tc(g, "Validate nút gửi", "Chỉ có dấu cách → nút gửi tắt", EMPTY, "Ô nhập: '     ' (5 dấu cách)", "1. Gõ 5 dấu cách\n2. Quan sát nút gửi",
   "1. Nút gửi vẫn mờ, không bấm được", "medium", A, "")
tc(g, "Validate nút gửi", "1 ký tự → nút gửi bật", EMPTY, "Ô nhập: 'a'", "1. Gõ 'a'\n2. Quan sát nút gửi",
   "1. Nút gửi bật", "medium", A, "")
tc(g, "Validate nút gửi", "Chỉ kèm ảnh, chưa gõ chữ → nút gửi tắt", EMPTY + "\n5. Đã cấp quyền Photos", "N/A", "1. Bấm dấu cộng\n2. Chọn 1 ảnh\n3. Quan sát nút gửi",
   "1. Ô thu nhỏ ảnh hiện trên ô nhập\n2. Nút gửi vẫn mờ, không bấm được", "high", A, "")
tc(g, "Validate độ dài", "Gõ đúng 1000 ký tự", EMPTY, "Chuỗi 1000 ký tự 'a'", "1. Dán chuỗi 1000 ký tự vào ô nhập",
   "1. Ô nhập nhận đủ 1000 ký tự\n2. Nút gửi bật", "medium", A, "Đếm độ dài value ô nhập")
tc(g, "Validate độ dài", "Dán 1001 ký tự → chặn ở 1000 + dòng nhắc", EMPTY, "Chuỗi 1001 ký tự 'a'", "1. Dán chuỗi 1001 ký tự vào ô nhập",
   "1. Ô nhập chỉ giữ 1000 ký tự\n2. Dòng nhắc dưới ô nhập 'Up to 1000 characters' (spec: 'Tối đa 1000 ký tự')", "high", A, "[⚠️ AMBIGUOUS C2] chữ spec tiếng Việt, app tiếng Anh")
tc(g, "UI ô nhập", "Ô nhập cao theo số dòng tới mức giới hạn", EMPTY, "8 dòng chữ (xuống dòng)", "1. Gõ 8 dòng chữ",
   "1. Ô nhập cao dần theo số dòng\n2. Tới mức giới hạn thì không cao thêm, chữ cuộn bên trong ô", "low", A, "So chiều cao ô nhập sau 1, 3, 8 dòng")
tc(g, "Logic dấu cộng", "Bấm dấu cộng mở thư viện có ảnh + video", EMPTY + "\n5. Đã cấp quyền Photos, thư viện có ảnh và video", "N/A", "1. Bấm nút dấu cộng",
   "1. Thư viện ảnh của máy mở\n2. Hiện cả ảnh và video", "high", A, "Picker hệ thống — PhotoPickerAutomator")
tc(g, "Logic dấu cộng", "Chọn ảnh → ô thu nhỏ có nút x", EMPTY + "\n5. Đã cấp quyền Photos", "N/A", "1. Bấm dấu cộng\n2. Chọn 1 ảnh",
   "1. Thư viện đóng\n2. Ô thu nhỏ ảnh hiện phía trên ô nhập, có nút x", "high", A, "")
tc(g, "Logic dấu cộng", "Chọn tệp khác thì thay tệp cũ", EMPTY + "\n5. Đã kèm ảnh A", "N/A", "1. Bấm dấu cộng\n2. Chọn ảnh B",
   "1. Chỉ còn 1 ô thu nhỏ — ảnh B", "medium", A, "Đếm ô thu nhỏ = 1, dấu vân màu đổi")
tc(g, "Logic dấu cộng", "Bấm x bỏ tệp đã kèm", EMPTY + "\n5. Đã kèm 1 ảnh", "N/A", "1. Bấm x trên ô thu nhỏ",
   "1. Ô thu nhỏ biến mất\n2. Chữ trong ô nhập giữ nguyên", "medium", A, "")
tc(g, "Validate video", "Video dài 40 giây không được kèm", EMPTY + "\n5. Thư viện có video 40 giây", "Video 40s", "1. Gõ 'abc'\n2. Bấm dấu cộng\n3. Chọn video 40 giây",
   "1. Không có ô thu nhỏ nào được thêm\n2. Thông báo 'Video must be 15 seconds or shorter' (spec: 'Video tối đa 15 giây')\n3. Ô nhập giữ nguyên 'abc'", "high", S, "Cần video 40s trong thư viện (PhotoLibrarySeeder)")
tc(g, "Validate video", "Video đúng 15 giây được kèm", EMPTY + "\n5. Thư viện có video 15 giây", "Video 15s", "1. Bấm dấu cộng\n2. Chọn video 15 giây",
   "1. Ô thu nhỏ video có nút phát nhỏ hiện trên ô nhập", "medium", S, "Cần video 15s trong thư viện")
tc(g, "Edge quyền", "Chưa cấp quyền Photos → hỏi quyền khi bấm dấu cộng", EMPTY + "\n5. Quyền Photos chưa hỏi (cài mới app)", "N/A", "1. Bấm dấu cộng",
   "1. Hộp thoại xin quyền thư viện ảnh của iOS hiện", "medium", S, "Cần gỡ/cài lại app trước case")
tc(g, "Edge quyền", "Từ chối quyền Photos → thông báo + nút mở Cài đặt", EMPTY + "\n5. Quyền Photos = None", "N/A", "1. Bấm dấu cộng",
   "1. Thông báo cần quyền kèm nút mở Cài đặt\n2. Bấm nút thì mở trang Settings của app", "medium", S, "Đặt quyền None qua SettingsPhotosDriver")
tc(g, "Logic gửi tạo ảnh", "Gửi yêu cầu tạo ảnh từ chữ", EMPTY, "Prompt: 'a cat wearing a wizard hat, watercolor style'", "1. Gõ prompt\n2. Bấm gửi",
   "1. Câu chào mất ngay\n2. Tin nhắn người dùng hiện bên phải trong bong bóng xám\n3. Ô đang tạo có biểu tượng lấp lánh + thanh tiến trình\n4. Nút gửi khoá tới khi có kết quả\n5. Kết quả ảnh về trong cuộc trò chuyện", "critical", A, "Tốn 1 lượt")
tc(g, "Logic gửi sửa ảnh", "Kèm 1 ảnh + yêu cầu sửa", EMPTY + "\n5. Đã cấp quyền Photos", "Prompt: 'turn into an oil painting'", "1. Bấm dấu cộng, chọn 1 ảnh\n2. Gõ prompt\n3. Bấm gửi",
   "1. Tin nhắn bên phải gồm ảnh thu nhỏ + câu yêu cầu\n2. Genius trả về 1 ảnh mới đã sửa\n3. Ảnh gốc trong Photos không đổi", "critical", A, "Tốn 1 lượt; đếm asset Photos trước/sau")
tc(g, "Logic gửi tạo video", "Yêu cầu video từ ảnh", EMPTY + "\n5. Đã cấp quyền Photos", "Prompt: 'make a video of this person waving'", "1. Kèm 1 ảnh\n2. Gõ prompt\n3. Bấm gửi",
   "1. Ô đang tạo dạng khung dọc, dòng 'Video takes a few minutes'\n2. Vẫn cuộn xem tin cũ được\n3. Xong: ô đổi thành video có nút phát ở giữa", "high", A, "Tốn 1 lượt, chờ ≤5 phút")
tc(g, "Logic gửi sửa video", "Kèm video ≤15s + yêu cầu sửa", EMPTY + "\n5. Thư viện có video ≤15s", "Prompt: 'make it look like an old film'", "1. Kèm video\n2. Gõ prompt\n3. Bấm gửi",
   "1. Tin nhắn bên phải có ô thu nhỏ video có nút phát + câu yêu cầu\n2. Ô đang tạo báo video mất vài phút\n3. Xong: video mới đã sửa; video gốc trong máy không đổi", "high", S, "Tốn 1 lượt; cần video ≤15s")
tc(g, "Edge rời màn", "Đổi tab trong lúc đang tạo, quay lại thấy kết quả", EMPTY, "Prompt: 'a red apple on a table'", "1. Gửi prompt\n2. Ngay lập tức chạm tab Home\n3. Chờ khay báo Ready\n4. Quay lại tab Chat",
   "1. Việc tạo chạy tiếp ở nền (khay trên thanh tab)\n2. Quay lại Chat thấy kết quả trong cuộc trò chuyện", "high", A, "Tốn 1 lượt")
tc(g, "Logic yêu cầu không để tạo", "Gửi câu chào", EMPTY, "Prompt: 'hello'", "1. Gửi 'hello'",
   "1. Genius trả lời 1 dòng chữ, không có bong bóng, gợi ý mô tả ảnh/video\n2. Không có hàng nút Lưu · Chia sẻ · Tạo lại\n3. Số lượt giảm 1", "medium", A, "[⚠️ AMBIGUOUS C8] chờ BE trừ lượt cho câu trả lời chữ")

# ---- CH03 Chuỗi chào bán & lượt (§4) ------------------------------------------------------
g = "AI Chat — Chuỗi chào bán & trừ lượt (236 · 238 · 237 · 481)"
tc(g, "Logic chặn người chưa mua gói", "Bấm gửi → paywall 236", NOSUB, "Prompt: 'a cat in a garden'", "1. Mở Chat\n2. Gõ prompt\n3. Bấm gửi",
   "1. Màn hình mua gói Chat (236) mở\n2. Chưa có tin nhắn nào được gửi, chưa trừ lượt", "critical", A, "Máy hiện không gói — chạy ngay")
tc(g, "Logic chặn người chưa mua gói", "Đóng 236 → last chance 238", NOSUB + "\n4. RC: show_238_lastchance_chat_paywall = true\n5. Paywall 236 đang mở", "N/A", "1. Bấm đóng paywall 236",
   "1. Màn last chance (238) mở thay chỗ 236", "high", A, "")
tc(g, "Logic chặn người chưa mua gói", "Đóng 238 → topup 237 (ví rỗng, hôm nay chưa chào)", NOSUB + "\n4. RC: enable_topup = true\n5. Hôm nay chưa chào topup\n6. Màn 238 đang mở", "N/A", "1. Bấm đóng màn 238",
   "1. Màn nạp lượt (237) mở", "high", A, "Chỉ đúng ở lần chào đầu tiên trong ngày")
tc(g, "Logic chặn người chưa mua gói", "Đóng hết chuỗi → về Chat, chữ còn nguyên", NOSUB + "\n4. Chuỗi 236 → 238 → 237 đang mở", "Prompt: 'a cat in a garden'", "1. Đóng lần lượt từng màn tới khi về Chat",
   "1. Về màn Chat\n2. Ô nhập còn nguyên 'a cat in a garden'\n3. Không có tin nhắn nào được gửi", "critical", A, "")
tc(g, "Logic chặn người chưa mua gói", "Topup chỉ chào 1 lần/ngày", NOSUB + "\n4. Hôm nay đã chào topup 237", "N/A", "1. Bấm gửi\n2. Đóng 236\n3. Đóng 238",
   "1. Không mở 237 lần nữa\n2. Về màn Chat, chữ còn nguyên", "medium", A, "Chạy ngay sau case trên")
tc(g, "Logic chặn người chưa mua gói", "Ví còn lượt topup → đóng chuỗi thì gửi luôn", BASE + "\n3. User chưa mua gói, ví topup ≥1", "Prompt: 'a cat in a garden'", "1. Bấm gửi\n2. Đóng 236\n3. Đóng 238",
   "1. Không chào 237\n2. Yêu cầu tự gửi đi, có kết quả\n3. Ví topup giảm 1", "high", S, "Cần mua topup sandbox (user nhập mật khẩu)")
tc(g, "Logic mua trong chuỗi", "Mua gói ở 236 → yêu cầu tự gửi", NOSUB + "\n4. Paywall 236 đang mở với prompt chờ", "N/A", "1. Mua gói thành công trên 236",
   "1. Paywall đóng\n2. Yêu cầu tự gửi đi, không phải bấm gửi lại\n3. Trừ 1 lượt", "high", S, "[⚠️ AMBIGUOUS C9] (giả định); cần user nhập mật khẩu sandbox")
tc(g, "Logic hết lượt", "Có gói nhưng hết lượt → hộp thoại 481", BASE + "\n3. User đã mua gói, lượt hôm nay = 0, ví topup = 0", "Prompt: 'a cat'", "1. Bấm gửi",
   "1. Hộp thoại hết lượt (481) hiện\n2. Số lượt không đổi", "critical", S, "Cần dùng hết lượt ngày (20/ngày)")
tc(g, "Logic hết lượt", "Đóng 481 → về Chat, chữ còn", BASE + "\n3. Hộp thoại 481 đang mở", "N/A", "1. Đóng hộp thoại 481",
   "1. Về màn Chat\n2. Chữ đang gõ còn nguyên\n3. Không có tin nhắn nào được gửi", "high", S, "")
tc(g, "Logic hết lượt", "481 → nút mua thêm mở màn nạp lượt", BASE + "\n3. Hộp thoại 481 đang mở", "N/A", "1. Bấm nút mua thêm trên 481",
   "1. Màn nạp thêm lượt mở", "medium", S, "")
tc(g, "Logic trừ lượt", "Tạo thành công trừ đúng 1 lượt", SUB, "Prompt: 'a blue car'", "1. Ghi số lượt ở huy hiệu Home\n2. Gửi prompt trong Chat, chờ kết quả\n3. Xem huy hiệu Home và My Balance",
   "1. Huy hiệu Home giảm đúng 1\n2. My Balance hiện cùng số với huy hiệu", "critical", A, "Tốn 1 lượt")
tc(g, "Logic trừ lượt", "Tạo video cũng chỉ trừ 1 lượt", SUB, "Prompt video từ ảnh", "1. Ghi số lượt\n2. Gửi yêu cầu tạo video, chờ xong\n3. Xem lại số lượt",
   "1. Số lượt giảm đúng 1", "high", A, "Gộp với case tạo video để tiết kiệm lượt")

# ---- CH04 Một lần gửi làm một việc (§4) ---------------------------------------------------
g = "AI Chat — Một lần gửi làm một việc (901)"
tc(g, "Logic tách việc", "Xin 2 ảnh cùng loại", EMPTY, "Prompt: 'create 2 images of a cat'", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Chỉ 1 ảnh, trừ 1 lượt\n2. Ô nhập điền sẵn 'create 1 images of a cat' (chưa gửi)\n3. Dưới kết quả: 'I do one thing at a time. The rest is ready in the box — tap send to continue.'\n4. Bong bóng người dùng giữ nguyên câu đã gõ", "high", A, "[⚠️ AMBIGUOUS C3] câu điền sẵn tiếng Anh tách theo luật chữ trên máy")
tc(g, "Logic tách việc", "Ảnh + video khác chủ thể", EMPTY, "Prompt: 'create 1 image of a cat and 1 video of a dog dancing'", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Chỉ 1 ảnh con mèo, trừ 1 lượt\n2. Ô nhập điền sẵn phần 'video con chó', chưa gửi\n3. Không có video nào được tạo", "high", A, "")
tc(g, "Logic tách việc", "Ảnh rồi video của chính ảnh đó", EMPTY, "Prompt: 'create an image of a cat then make a video of it running'", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Tạo ảnh trước, trừ 1 lượt\n2. Ô nhập điền sẵn 'make a video of it running'\n3. Thanh 'Editing this' hiện với ảnh vừa tạo", "high", A, "")
tc(g, "Logic tách việc", "Đang gõ chữ khi kết quả về → không ghi đè", EMPTY, "Prompt gộp + chữ gõ thêm 'xyz'", "1. Gửi prompt gộp\n2. Trong lúc chờ gõ 'xyz'\n3. Chờ kết quả",
   "1. Ô nhập vẫn là 'xyz'\n2. Dưới kết quả: 'I do one thing at a time. Send the rest as a new message.'", "medium", A, "")
tc(g, "Logic tách việc", "Câu > 160 ký tự không bị tách", EMPTY, "Prompt 170 ký tự có 'and' nhắc ảnh + video", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Chỉ 1 kết quả, trừ 1 lượt\n2. Ô nhập để trống\n3. Dòng 'I do one thing at a time. Send the rest as a new message.'", "low", A, "")

# ---- CH05 Lỗi & mất mạng (§3 · §4) --------------------------------------------------------
g = "AI Chat — Lỗi, quá thời gian & mất mạng (901)"
tc(g, "Edge mạng", "Bấm gửi khi mất mạng", SUB + "\n4. Bật Airplane Mode + tắt Wi-Fi", "Prompt: 'a cat'", "1. Gõ prompt\n2. Bấm gửi",
   "1. Thông báo mất mạng 'No internet connection. Please try again.' kèm nút Thử lại (spec: 'Không có kết nối mạng')\n2. Số lượt không đổi", "critical", S, "")
tc(g, "Edge mạng", "Mất mạng giữa lúc đang tạo", SUB, "Prompt: 'a cat'", "1. Gửi prompt\n2. Khi ô đang tạo hiện: bật Airplane Mode 60s\n3. Tắt Airplane Mode",
   "1. Ô đang tạo đổi thành dòng lỗi kèm nút Thử lại, hoặc lượt hoàn tất khi có mạng lại\n2. Lượt lỗi không bị trừ", "high", S, "[⚠️ AMBIGUOUS C6] spec không định nghĩa mất mạng giữa chừng — dùng tiêu chí như TC-BG01-015")
tc(g, "Edge tạo lỗi", "Bấm Thử lại trên ô lỗi gửi lại đúng yêu cầu", SUB + "\n4. Có một ô lỗi kèm nút Thử lại", "N/A", "1. Có mạng\n2. Bấm Thử lại",
   "1. Gửi lại đúng yêu cầu cũ\n2. Ra kết quả, trừ 1 lượt", "high", S, "Dựng ô lỗi bằng case mất mạng giữa chừng")
tc(g, "Edge quá thời gian", "Quá 180 giây chưa có kết quả", SUB, "N/A", "1. Gửi yêu cầu\n2. Chờ quá 180 giây không có kết quả",
   "1. Ô đang tạo đổi thành 'This is taking longer than expected.' kèm Thử lại\n2. Không trừ lượt", "medium", S, "Không ép được server chậm (không mock) — chỉ chạy được khi gặp thật; [⚠️ AMBIGUOUS C4] kết quả về muộn có trừ lượt?")
tc(g, "Edge nội dung", "Yêu cầu vi phạm nội dung bị từ chối", SUB, "Prompt vi phạm chính sách (bạo lực)", "1. Gửi prompt",
   "1. Genius trả lời 1 dòng chữ từ chối lịch sự\n2. Không có ảnh/video\n3. Trừ 1 lượt", "medium", A, "Dùng prompt bạo lực nhẹ đã thống nhất với QA lead")

# ---- CH06 Kết quả & thao tác (§5) ---------------------------------------------------------
g = "AI Chat — Kết quả & thao tác (901 → 902)"
tc(g, "UI kết quả", "Bố cục câu trả lời có kết quả", HASRES, "N/A", "1. Quan sát câu trả lời có kết quả gần nhất",
   "1. Không có bong bóng\n2. Một dòng chữ ngắn, dưới là ảnh đúng tỉ lệ đã tạo\n3. Dưới cùng hàng nút Lưu · Chia sẻ · Tạo lại", "high", A, "[⚠️ AMBIGUOUS C5] hàng nút chờ design/CPO chốt")
tc(g, "UI kết quả", "Câu trả lời chỉ có chữ không có hàng nút", SUB + "\n4. Cuộc trò chuyện có câu trả lời chỉ có chữ", "N/A", "1. Quan sát câu trả lời chữ",
   "1. Không có hàng nút Lưu · Chia sẻ · Tạo lại", "medium", A, "")
tc(g, "UI kết quả", "Nút báo cáo trên mỗi kết quả", HASRES, "N/A", "1. Quan sát kết quả",
   "1. Có nút Report (cờ) trên kết quả", "medium", A, "")
tc(g, "Logic xem kết quả", "Bấm kết quả mở 902 phủ kín", HASRES, "N/A", "1. Bấm vào ảnh kết quả",
   "1. Màn Xem kết quả (902) phủ kín màn hình, ẩn thanh tab\n2. Có nút đóng, Lưu, Chia sẻ", "high", A, "")
tc(g, "Logic xem kết quả", "Vuốt xuống đóng 902 về đúng chỗ", HASRES, "N/A", "1. Mở 902 từ kết quả\n2. Kéo ảnh xuống quá 1/4 chiều cao rồi thả",
   "1. Màn xem đóng\n2. Về đúng vị trí đang xem trong Chat", "medium", A, "")
tc(g, "Logic xem kết quả", "Kéo chưa tới 1/4 thì ảnh về giữa", HASRES, "N/A", "1. Mở 902\n2. Kéo ảnh xuống 1/8 chiều cao rồi thả",
   "1. Màn xem không đóng\n2. Ảnh trở về giữa màn", "low", A, "")
tc(g, "Logic lưu", "Bấm Lưu dưới kết quả", HASRES + "\n5. Đã cấp quyền Photos", "N/A", "1. Bấm Lưu dưới kết quả",
   "1. Thông báo 'Saved' ở đáy ~2 giây (spec: 'Đã lưu')\n2. Thư viện Photos có thêm 1 tệp", "high", A, "Đếm PHAsset trước/sau")
tc(g, "Logic chia sẻ", "Bấm Chia sẻ mở bảng chia sẻ", HASRES, "N/A", "1. Bấm Chia sẻ dưới kết quả",
   "1. Bảng chia sẻ của iOS mở với đúng tệp đó", "medium", A, "Đóng share sheet, không bấm app nhắn tin")
tc(g, "Logic tạo lại", "Bấm Tạo lại trừ lượt, thêm kết quả mới", HASRES, "N/A", "1. Ghi số lượt\n2. Bấm Tạo lại dưới kết quả\n3. Chờ kết quả",
   "1. Có câu trả lời mới bên dưới, kết quả cũ vẫn còn\n2. Số lượt giảm 1", "high", A, "Tốn 1 lượt")
tc(g, "Logic Library", "Kết quả Chat hiện đầu tab Library", HASRES, "N/A", "1. Mở tab Library",
   "1. Kết quả mới nhất của Chat nằm đầu lưới\n2. Video có nút phát nhỏ giữa ô", "high", A, "So dấu vân ô đầu Library với kết quả")
tc(g, "Logic Library", "Xoá kết quả ở Library → Chat hiện 'Tệp đã bị xoá'", HASRES, "N/A", "1. Xoá kết quả mới nhất trong tab Library\n2. Mở tab Chat",
   "1. Ô kết quả đó trong Chat hiện dòng báo tệp đã bị xoá", "medium", A, "")
tc(g, "Edge mạng 902", "Kết quả chưa tải về + mất mạng → khung xám + Thử lại", HASRES + "\n5. Kết quả chưa từng tải về máy (cài lại app)\n6. Mất mạng", "N/A", "1. Bấm vào kết quả",
   "1. 902 hiện khung xám 'Không tải được' kèm nút Thử lại\n2. Có mạng lại bấm Thử lại thì hiện kết quả", "low", S, "Dựng 'chưa tải về' khó — cài lại app + Restore")

# ---- CH07 Sửa tiếp & vuốt chọn (§6) -------------------------------------------------------
g = "AI Chat — Sửa tiếp & vuốt chọn kết quả (901)"
tc(g, "Logic sửa tiếp", "Gửi tiếp không chọn gì → sửa kết quả gần nhất", HASRES, "Prompt: 'make the background a beach'", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Kết quả mới là bản sửa của kết quả ngay trước\n2. Kết quả cũ vẫn còn\n3. Trừ 1 lượt", "high", A, "Tốn 1 lượt; so dấu vân")
tc(g, "Logic vuốt chọn", "Vuốt kết quả cũ qua ngưỡng 1/5 → thanh Editing this", TWORES, "N/A", "1. Kéo kết quả thứ 2 từ dưới lên sang phải ~1/3 chiều rộng\n2. Thả tay",
   "1. Kết quả trở về chỗ cũ\n2. Trên thanh soạn hiện thanh 'Editing this' + ảnh thu nhỏ của đúng kết quả đó + nút x\n3. Bàn phím mở", "critical", A, "Rung nhẹ không đo được bằng automation")
tc(g, "Logic vuốt chọn", "Kéo chưa tới 1/5 → không chọn", TWORES, "N/A", "1. Kéo kết quả sang phải ~1/10 chiều rộng\n2. Thả tay",
   "1. Không có thanh 'Editing this'", "high", A, "")
tc(g, "Logic vuốt chọn", "Gửi yêu cầu cho kết quả đã chọn", TWORES + "\n5. Thanh 'Editing this' đang chọn kết quả cũ", "Prompt: 'make it black and white'", "1. Gửi prompt\n2. Chờ kết quả",
   "1. Tin nhắn người dùng có ảnh thu nhỏ của kết quả đã chọn phía trên câu chữ\n2. Kết quả mới là bản sửa của kết quả đã chọn, không phải kết quả gần nhất\n3. Thanh 'Editing this' biến mất", "critical", A, "Tốn 1 lượt")
tc(g, "Logic vuốt chọn", "Bấm x bỏ chọn, chữ còn nguyên", TWORES + "\n5. Thanh 'Editing this' đang hiện, ô nhập có 'abc'", "N/A", "1. Bấm x trên thanh 'Editing this'",
   "1. Thanh biến mất\n2. Ô nhập vẫn 'abc'", "medium", A, "")
tc(g, "Logic vuốt chọn", "Vuốt kết quả khác → đổi đối tượng", TWORES + "\n5. Đang chọn kết quả A", "N/A", "1. Vuốt kết quả B sang phải qua ngưỡng",
   "1. Thanh 'Editing this' đổi sang ảnh thu nhỏ của B\n2. Chỉ 1 kết quả được chọn", "medium", A, "")
tc(g, "Logic vuốt chọn", "Đang chọn kết quả mà kèm ảnh → thanh chọn mất", TWORES + "\n5. Đang chọn 1 kết quả", "N/A", "1. Bấm dấu cộng\n2. Chọn 1 ảnh",
   "1. Thanh 'Editing this' biến mất\n2. Chỉ còn ô thu nhỏ ảnh vừa chọn", "medium", A, "")
tc(g, "Edge vuốt", "Vuốt tin nhắn người dùng / câu trả lời chữ → không chọn", SUB + "\n4. Có tin nhắn người dùng và câu trả lời chữ", "N/A", "1. Vuốt tin nhắn người dùng sang phải\n2. Vuốt câu trả lời chữ sang phải",
   "1. Tin nhắn không di chuyển, không hiện mũi tên\n2. Không có thanh 'Editing this'", "medium", A, "")
tc(g, "Edge chưa có kết quả", "Chưa có kết quả → tạo mới hoàn toàn", EMPTY, "Prompt: 'a lighthouse at night'", "1. Gửi prompt",
   "1. Có kết quả mới, không có thanh 'Editing this'", "low", A, "Gộp với case gửi tạo ảnh")
tc(g, "Edge tệp đã xoá", "Kết quả đã chọn bị xoá ở Library → 'Can't find the image to edit'", TWORES + "\n5. Chọn kết quả X, rồi xoá X ở Library", "Prompt: 'add a hat'", "1. Quay lại Chat\n2. Gõ prompt, bấm gửi",
   "1. Thông báo 'Can't find the image to edit' (spec: 'Không tìm thấy ảnh cần sửa')\n2. Thanh 'Editing this' biến mất, chữ còn\n3. Số lượt không đổi", "medium", A, "")

# ---- CH08 Ô nhập kính & sửa trên 902 (§7) -------------------------------------------------
g = "AI Chat — Ô nhập kính & sửa trên màn xem (901 · 902)"
tc(g, "UI ô nhập kính", "Ô nhập là lớp kính, chữ vẫn rõ", HASRES, "N/A", "1. Cuộn để kết quả nằm ngay sau ô nhập\n2. Quan sát ô nhập",
   "1. Thấy lờ mờ nội dung phía sau ô nhập\n2. Chữ, nút dấu cộng, nút gửi vẫn rõ", "low", A, "So ảnh vùng ô nhập khi nền sáng/tối khác nhau")
tc(g, "UI ô nhập kính", "Giảm độ trong suốt → nền đặc", HASRES + "\n5. Bật Reduce Transparency trong Trợ năng", "N/A", "1. Mở Chat\n2. Quan sát ô nhập",
   "1. Ô nhập nền đặc, không nhìn xuyên được", "low", S, "Bật Reduce Transparency qua app Settings")
tc(g, "UI 902 từ Chat", "902 mở từ Chat có ô nhập ở đáy", HASRES, "N/A", "1. Bấm kết quả trong Chat",
   "1. Ô nhập kính ở đáy\n2. Không có nút dấu cộng, không có thanh 'Editing this'\n3. Ảnh và nút Report nằm trọn phía trên ô nhập", "high", A, "So frame ảnh với frame ô nhập")
tc(g, "UI 902 từ Library", "902 mở từ Library không có ô nhập", HASRES, "N/A", "1. Mở tab Library\n2. Bấm kết quả",
   "1. Màn xem không có ô nhập", "high", A, "Đã thấy ở suite Library (không có Edit More)")
tc(g, "Logic sửa trên 902", "Gõ + gửi trên 902 sửa đúng kết quả đang xem", TWORES, "Prompt: 'add sunglasses'", "1. Mở kết quả cũ (không phải mới nhất)\n2. Gõ prompt, bấm gửi",
   "1. Màn xem đóng\n2. Khung chat hiện lượt đang tạo\n3. Kết quả mới là bản sửa của đúng kết quả vừa xem\n4. Trừ 1 lượt", "critical", A, "Tốn 1 lượt")
tc(g, "Validate 902", "Ô nhập rỗng trên 902 → nút gửi tắt", HASRES, "N/A", "1. Mở 902 từ Chat\n2. Quan sát nút gửi",
   "1. Nút gửi tắt\n2. Không có lượt tạo mới", "medium", A, "")
tc(g, "Logic chặn trên 902", "Chưa mua gói gửi từ 902 → paywall", NOSUB + "\n4. Cuộc trò chuyện có kết quả cũ", "Prompt: 'add a hat'", "1. Mở kết quả\n2. Gõ prompt, bấm gửi",
   "1. Paywall 236 mở\n2. Không trừ lượt", "high", A, "Máy hiện không gói — chạy được nếu cuộc trò chuyện cũ còn kết quả")
tc(g, "Edge mạng 902", "Mất mạng khi gửi từ 902", HASRES + "\n5. Mất mạng", "Prompt: 'add a hat'", "1. Mở kết quả\n2. Gõ prompt, bấm gửi",
   "1. Màn xem đóng, hiện thông báo mất mạng\n2. Ô soạn ở khung chat còn 'add a hat' và kết quả đó vẫn là đối tượng sửa\n3. Không trừ lượt", "high", S, "")
tc(g, "Logic giữ phần soạn", "Đóng 902 không gửi → ô soạn như cũ", TWORES + "\n5. Ô soạn có ảnh kèm + đang chọn kết quả A", "N/A", "1. Mở kết quả B\n2. Đóng màn xem không gửi",
   "1. Ô soạn vẫn có ảnh kèm\n2. Kết quả A vẫn được chọn", "medium", A, "")
tc(g, "Logic giữ phần soạn", "Gửi từ 902 → ảnh kèm trở lại, kết quả chọn trước bị bỏ", TWORES + "\n5. Ô soạn có ảnh kèm", "Prompt: 'add a hat'", "1. Mở kết quả B\n2. Gõ prompt, gửi từ 902",
   "1. Lượt gửi dùng kết quả B\n2. Ảnh kèm trở lại ô soạn khung chat\n3. Không có kết quả nào được chọn lại", "medium", A, "Tốn 1 lượt")
tc(g, "Logic giữ phần soạn", "Kéo xuống đóng 902 khi đã gõ chữ", HASRES, "Chữ: 'draft text'", "1. Mở 902\n2. Gõ 'draft text'\n3. Kéo ảnh xuống quá 1/4 rồi thả",
   "1. Màn xem đóng\n2. 'draft text' còn trong ô soạn khung chat", "medium", A, "")
tc(g, "UI 902 bàn phím", "Ô nhập nằm trên bàn phím", HASRES, "N/A", "1. Mở 902\n2. Chạm ô nhập",
   "1. Ô nhập nằm ngay trên bàn phím, không bị che", "low", A, "So frame ô nhập với frame bàn phím")

# ---- VB01 Dải Try the newest vibes (§8) ---------------------------------------------------
g = "AI Chat — Dải 'Try the newest vibes' ở Home (301)"
tc(g, "UI dải vibes", "Dải đứng đầu tab For You", VIBES, "N/A", "1. Mở Home, tab For You\n2. Quan sát dải ngay dưới thanh chọn",
   "1. Dải đầu tiên là 'Try the newest vibes' (hoặc tên category vibes)\n2. Các ô nằm ngang, mỗi ô 1 tên phong cách trên ảnh bìa riêng", "high", A, "")
tc(g, "UI dải vibes", "Chưa có category vibes → 7 phong cách mặc định", VIBES + "\n5. CMS chưa dựng category 'vibes'", "N/A", "1. Mở Home, cuộn ngang dải vibes",
   "1. Đủ 7 ô: Cinematic · Polaroid · Black and White · Retro · Golden Hour · Neon Night · Oil Painting", "high", A, "Spec: category chưa được tạo — kiểm thực tế trên máy")
tc(g, "Logic dải vibes", "Phong cách vibes không hiện trùng ở chỗ khác", VIBES, "N/A", "1. Lướt tab For You\n2. Mở tab Trending, New\n3. Mở See All",
   "1. Không thấy các phong cách vibes ngoài dải 'Try the newest vibes'", "medium", A, "")
tc(g, "Logic dải vibes", "Chạm ô phong cách mở thư viện ảnh, chưa sang Chat", VIBES, "N/A", "1. Chạm ô 'Polaroid'",
   "1. Thư viện ảnh của máy mở ngay\n2. Không thấy khung trò chuyện", "critical", A, "")
tc(g, "Edge mạng vibes", "Mất mạng khi mở Home → dải vẫn hiện", VIBES + "\n5. Mất mạng", "N/A", "1. Mở lại app tới Home",
   "1. Dải vibes hiện theo dữ liệu đã lưu, chưa từng tải thì bộ mặc định", "medium", S, "")

# ---- VB02 Chọn ảnh & xem lại (§9 · §10) ---------------------------------------------------
g = "AI Chat — Chọn ảnh & xem lại của dải vibes"
tc(g, "Logic chọn ảnh", "Chọn 1 ảnh → màn xem lại đúng ảnh", VIBES, "N/A", "1. Chạm ô 'Retro'\n2. Chọn 1 ảnh",
   "1. Thư viện đóng\n2. Màn xem lại hiện đúng ảnh vừa chọn", "critical", A, "")
tc(g, "UI xem lại", "Màn xem lại có ảnh, tên phong cách, nút đi tiếp bật", VIBES + "\n5. Đang ở màn xem lại", "N/A", "1. Quan sát màn",
   "1. Ảnh vừa chọn\n2. Tên phong cách đã chạm\n3. Nút đi tiếp đang bật\n4. Nút 'Change photo'", "high", A, "")
tc(g, "Logic chọn ảnh", "Huỷ thư viện → về Home", VIBES, "N/A", "1. Chạm ô phong cách\n2. Bấm Huỷ trong thư viện",
   "1. Về Home, dải vibes vẫn ở chỗ cũ\n2. Không mở màn xem lại, không mở Chat", "high", A, "")
tc(g, "Logic xem lại", "Change photo → thay ảnh", VIBES + "\n5. Đang ở màn xem lại với ảnh A", "N/A", "1. Bấm 'Change photo'\n2. Chọn ảnh B",
   "1. Màn xem lại đổi sang ảnh B", "medium", A, "")
tc(g, "Edge quyền", "Từ chối quyền Photos khi chạm ô vibes", BASE + "\n3. RC: enable_home_vibes_feature = true\n4. Quyền Photos = None", "N/A", "1. Chạm ô phong cách",
   "1. Lời nhắc kèm nút mở Cài đặt\n2. Người dùng ở lại Home\n3. Bấm nút thì mở trang Settings của app", "medium", S, "Đặt quyền None qua SettingsPhotosDriver")
tc(g, "Edge xem lại", "Mất mạng khi bấm đi tiếp → lỗi + thử lại", VIBES + "\n5. Đang ở màn xem lại\n6. Mất mạng", "N/A", "1. Bấm đi tiếp\n2. Bật lại mạng\n3. Bấm thử lại",
   "1. Thông báo lỗi kèm nút thử lại, ảnh giữ nguyên\n2. Sau khi có mạng, thử lại thì sang Chat", "medium", S, "")

# ---- VB03 Chat điền sẵn (§11) -------------------------------------------------------------
g = "AI Chat — Chat điền sẵn từ dải vibes (901)"
tc(g, "Logic điền sẵn", "Đi tiếp → Chat có ảnh + chữ mô tả, nút gửi bật, chưa gửi", VIBES + "\n5. Đang ở màn xem lại", "N/A", "1. Bấm đi tiếp",
   "1. Mở Chat\n2. Ô soạn có ảnh vừa chọn + phần chữ mô tả của phong cách\n3. Nút gửi bật\n4. Chưa có tin nhắn nào được gửi", "critical", A, "")
tc(g, "Logic điền sẵn", "Gỡ ảnh → vẫn gửi được như tạo từ chữ", VIBES + "\n5. Chat đang điền sẵn ảnh + chữ", "N/A", "1. Bấm x gỡ ảnh\n2. Quan sát nút gửi",
   "1. Ảnh biến mất, chữ mô tả giữ nguyên\n2. Nút gửi vẫn bật\n3. Không mở lại thư viện ảnh", "high", A, "Không bấm gửi (chỉ kiểm trạng thái) nếu máy không gói")
tc(g, "Logic chặn", "Chưa mua gói bấm gửi điền sẵn → paywall, đóng thì còn nguyên", NOSUB + "\n4. Chat đang điền sẵn ảnh + chữ từ dải vibes", "N/A", "1. Bấm gửi\n2. Đóng hết chuỗi paywall",
   "1. Màn hình mua gói mở\n2. Về Chat: ảnh và chữ còn nguyên\n3. Bấm gửi lại thì lại hiện màn mua gói", "high", A, "Máy hiện không gói — chạy ngay")
tc(g, "Logic điền sẵn", "Gửi điền sẵn khi có gói → tạo ảnh", VIBES + "\n5. User đã mua gói, còn ≥10 lượt\n6. Chat đang điền sẵn", "N/A", "1. Bấm gửi\n2. Chờ kết quả",
   "1. Tin nhắn bên phải có ảnh + chữ mô tả\n2. Ra kết quả ảnh theo phong cách\n3. Trừ 1 lượt", "critical", A, "Tốn 1 lượt")

# ---- REG -----------------------------------------------------------------------------------
g = "Regression / E2E AI Chat"
tc(g, "E2E", "Chat → tạo → sửa tiếp → xem 902 → Lưu → Library", EMPTY, "Prompt 1: 'a cat in a garden' · Prompt 2: 'make it night'", "1. Gửi prompt 1, chờ kết quả\n2. Gửi prompt 2, chờ kết quả\n3. Mở kết quả mới, bấm Lưu\n4. Mở Library",
   "1. Hai kết quả trong Chat, kết quả 2 là bản sửa của kết quả 1\n2. Lưu có thông báo 'Saved'\n3. Library có 2 kết quả mới ở đầu", "critical", A, "Tốn 2 lượt")
tc(g, "Regression", "Kết quả Chat không làm hỏng luồng style Home", HASRES, "N/A", "1. Mở Home\n2. Chạy luồng i2i 1 ảnh tới Result",
   "1. Luồng i2i ra Result như trước\n2. Kết quả i2i vào Library", "high", A, "Tốn 1 lượt")

# ---- export -------------------------------------------------------------------------------
HDR = ['N°', 'Feature', 'Test Description', 'Sub-scenario', 'Precondition', 'Test Data', 'Action', 'Expected Result', 'PASS/FAIL', 'TC ID', 'Priority', 'Automation', 'Ghi chú automation']
WID = [5, 22, 24, 30, 42, 28, 42, 52, 11, 13, 10, 11, 30]
CODES = {k: f"CH{i:02d}" for i, k in enumerate([x for x in G if x.startswith("AI Chat —") and "vibes" not in x.lower() and "dải" not in x.lower()], 1)}
for i, k in enumerate([x for x in G if "vibes" in x.lower()], 1): CODES[k] = f"VB{i:02d}"
CODES["Regression / E2E AI Chat"] = "REG02"

wb = Workbook(); ws = wb.active; ws.title = "TC_AI_Chat"
thin = Side(style="thin", color="BBBBBB"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("B1:H1"); ws["B1"] = "spec: Genius 2.9.0 — IIP055.1.20 AI Chat (pageId 323393416 §3–11) — scope: AI Chat + dải Try the newest vibes"
ws["B1"].font = Font(name="Arial", size=11, bold=True)
for c, (h, w) in enumerate(zip(HDR, WID), 1):
    cell = ws.cell(2, c, h); cell.font = Font(name="Arial", size=10, bold=True)
    cell.fill = PatternFill("solid", fgColor="B6D7A8"); cell.border = border; cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[cell.column_letter].width = w
ws.freeze_panes = "A3"
row, n, spec, summary = 3, 0, {}, []
for group, items in G.items():
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
    sep = ws.cell(row, 1, f"▌ {group}"); sep.font = Font(name="Arial", size=10, bold=True); sep.fill = PatternFill("solid", fgColor="FFF2CC"); row += 1
    start, code, counts = row, CODES[group], {"Auto": 0, "Auto + setup": 0, "critical": 0, "high": 0, "medium": 0, "low": 0}
    prev_desc, desc_start = None, row
    for k, (desc, sub, pre, data, action, exp, prio, auto, note) in enumerate(items, 1):
        n += 1; tid = f"TC-{code}-{k:03d}"
        vals = [n, group if row == start else None, desc, sub, pre, data, action, exp, None, tid, prio, auto, note]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v); cell.alignment = wrap; cell.border = border; cell.font = Font(name="Arial", size=10)
        if desc != prev_desc:
            if prev_desc is not None and row - 1 > desc_start: ws.merge_cells(start_row=desc_start, start_column=3, end_row=row - 1, end_column=3)
            desc_start, prev_desc = row, desc
        else: ws.cell(row, 3).value = None
        counts[auto] += 1; counts[prio] += 1
        spec[tid] = {"title": f"{desc} — {sub}", "pre": pre, "data": data, "action": action, "expected": exp, "priority": prio, "feature": group}
        row += 1
    if row - 1 > desc_start: ws.merge_cells(start_row=desc_start, start_column=3, end_row=row - 1, end_column=3)
    if row - 1 > start: ws.merge_cells(start_row=start, start_column=2, end_row=row - 1, end_column=2)
    ws.cell(start, 2).alignment = Alignment(wrap_text=True, vertical="center")
    summary.append((group, len(items), counts["Auto"], counts["Auto + setup"], counts["critical"], counts["high"], counts["medium"], counts["low"]))

amb = wb.create_sheet("Ambiguity"); amb.append(["#", "Mục spec", "Điểm mơ hồ — cần PO/BA confirm", "TC liên quan"])
for r in [
    ("C1", "§3, §4, §5, §6", "Spec viết thông báo bằng tiếng Việt ('Cuộc trò chuyện mới', 'Không có kết nối mạng', 'Đã lưu', 'Video tối đa 15 giây', 'Tối đa 1000 ký tự', 'Không tìm thấy ảnh cần sửa'), app chỉ có chuỗi tiếng Anh. TC chấm theo nghĩa, chữ ghi theo app.", "CH01, CH02, CH05, CH06, CH07"),
    ("C2", "§4 Yêu cầu quá dài", "Dòng nhắc 'Tối đa 1000 ký tự' hiện lúc nào: chỉ khi chạm 1000 hay khi dán quá? App: 'Up to %d characters'.", "Validate độ dài"),
    ("C3", "§4 Một lần gửi một việc", "Câu điền sẵn với prompt tiếng Anh ('create 2 images' → 'create 1 images'?) — spec chỉ ví dụ tiếng Việt; luật tách câu chỉ nhận và · and · rồi · then.", "CH04"),
    ("C4", "§4 Quá thời gian chờ", "Kết quả về muộn sau khi đã báo quá 180 s vẫn lưu và bị trừ lượt (giả định) — PO xác nhận.", "Quá 180 giây"),
    ("C5", "§5 Hàng nút", "Hàng nút Lưu · Chia sẻ · Tạo lại dưới mỗi kết quả: bản thiết kế chỉ vẽ trạng thái trống — chờ design/CPO chốt.", "UI kết quả"),
    ("C6", "§4 · §13", "Mất mạng giữa lúc Chat đang tạo: spec không định nghĩa. Đề xuất dùng tiêu chí như BG01-015 (hoàn tất hoặc lỗi có Thử lại).", "Mất mạng giữa lúc tạo"),
    ("C7", "§11 vs §4", "§11 nói tạo lỗi/quá thời gian thì 'lượt dùng được hoàn lại', §4 nói 'không trừ lượt'. Cùng ý hay có trừ rồi hoàn?", "CH05, VB03"),
    ("C8", "§4 Câu chào", "Câu trả lời chỉ có chữ trừ 1 lượt — chờ BE có đường trừ lượt. Hiện BE chỉ trừ khi tạo ảnh/video.", "Gửi câu chào"),
    ("C9", "§4 Mua trên 236", "Mua gói thành công trên 236 thì yêu cầu tự gửi (giả định).", "Mua gói ở 236"),
    ("C10", "§11 hết lượt", "§11 'hiện lời mời mua thêm lượt' — là hộp thoại 481 như §4 hay màn topup?", "VB03"),
    ("C11", "RC", "enable_ai_chat_feature = false / enable_home_vibes_feature = false không sinh TC chạy: user chốt không đổi Remote Config khi chạy automation.", "—"),
]: amb.append(r)
setup = wb.create_sheet("Automation Setup"); setup.append(["Hạng mục", "Ghi chú"])
for r in [
    ("Platform", "iPhone 12 (iOS 26.5.2) đã pair; XCUITest, build bằng Xcode 26.6 (DEVELOPER_DIR)."),
    ("Chọn element", "Màn Chat không có accessibility identifier — chọn theo nhãn tiếng Anh (Chat, What's on your mind?, New conversation, Editing this, Regenerate, Report…)."),
    ("Trạng thái gói", "Nhóm 'Chưa mua gói' chạy khi máy không có gói (hiện tại). Nhóm 'Đã mua gói' cần mua lại Yearly sandbox — tester nhập mật khẩu."),
    ("Lượt tạo", "Mỗi lần gửi trừ 1 lượt, 20 lượt/ngày. Gom case để 1 lượt gửi phục vụ nhiều TC."),
    ("Mất mạng", "Airplane Mode + tắt Wi-Fi qua app Settings (SettingsNetworkDriver), luôn trả mạng ở defer."),
    ("Dữ liệu ảnh/video", "Thư viện cần ảnh mặt người, 1 video ≤15 s và 1 video 40 s (PhotoLibrarySeeder tạo được ảnh; video cần chép sẵn)."),
    ("Không làm", "Không đổi Remote Config, không mock/proxy server; case cần ép lỗi server/timeout chỉ chấm khi gặp thật."),
]: setup.append(r)
sm = wb.create_sheet("Summary"); sm.append(["Feature", "Số TC", "Auto", "Auto + setup", "critical", "high", "medium", "low"])
for r in summary: sm.append(r)
last = len(summary) + 1
sm.append(["Tổng"] + [f"=SUM({c}2:{c}{last})" for c in "BCDEFGH"])
for w in (amb, setup, sm):
    for col in w.columns: w.column_dimensions[col[0].column_letter].width = 40 if col[0].column_letter != "A" else 30
    for c in w[1]: c.font = Font(name="Arial", size=10, bold=True); c.fill = PatternFill("solid", fgColor="B6D7A8")
wb.save(OUT)
json.dump({"ui": spec}, open(JSON_OUT, "w"), ensure_ascii=False, indent=1)
print("TC:", n, {k: v for k, v in CODES.items()})
for r in summary: print(r)
