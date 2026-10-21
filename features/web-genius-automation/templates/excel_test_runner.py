"""Chay pytest theo file testcase Excel.

Sua so voi ban cu:
1. FEATURE_TARGETS dung dung ten trong cot Feature cua Excel
   (ban cu dung "Image Generator (/agent/image)" trong khi Excel ghi
   "Image Generator" -> 10/12 nhom bi bo qua am tham).
2. Feature khong co mapping -> BAO LOI va tra exit code != 0,
   thay vi return 0 khien CI bao xanh gia.
3. Them --tc de chay dung mot testcase theo so N° trong Excel.
4. Them --write-results de do ket qua nguoc vao cot PASS/FAIL.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_EXCEL = PROJECT_ROOT / "TC_Genius_Web_Automation.xlsx"
SHEET_NAME = "TC_Genius_Web"

# Khoa phai TRUNG KHOP voi gia tri cot Feature trong Excel.
FEATURE_TARGETS: dict[str, list[str]] = {
    "Authentication & Session": ["tests/test_auth.py"],
    "Home": ["tests/test_home.py"],
    "Image Generator": ["tests/test_image.py"],
    "Video Generator": ["tests/test_video.py"],
    "Voice Generator (Speech)": ["tests/test_tools_ui.py::test_tc18_voice_layout"],
    "Music Generator": ["tests/test_tools_ui.py::test_tc19_music_layout"],
    "AI Chat": ["tests/test_tools_ui.py::test_tc20_chat_layout"],
    "Projects": ["tests/test_projects.py"],
    "Daily Credit Quota": ["tests/test_credit.py"],
    "Navigation & Coming Soon": ["tests/test_navigation.py"],
    "Settings / Theme / Language": [
        "tests/test_navigation.py::test_tc34_dark_mode"
    ],
    "Regression & E2E": ["tests/test_regression.py"],
}


def load_cases_from_excel(excel_path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(Path(excel_path), data_only=True)
    sheet = (
        workbook[SHEET_NAME]
        if SHEET_NAME in workbook.sheetnames
        else workbook.active
    )

    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    headers: list[str] = []

    for idx, row in enumerate(rows):
        if row and row[0] == "N°":
            header_index = idx
            headers = [
                str(v).strip() if v is not None else "" for v in row
            ]
            break

    if header_index is None:
        raise ValueError(f"Khong tim thay dong header trong {excel_path}")

    cases: list[dict[str, Any]] = []
    current_feature = ""

    for row_number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        first = str(row[0]).strip() if row[0] is not None else ""

        # Dong phan cach nhom: "▌ IMAGE GENERATOR (/agent/image)"
        if first.startswith("▌"):
            continue

        numeric = first.replace(",", ".")
        if not numeric.replace(".", "", 1).isdigit():
            continue

        case: dict[str, Any] = {
            "n": int(float(numeric)),
            "row": row_number,
        }
        for key, value in zip(headers, row):
            if key:
                case[key] = value

        # Cot Feature duoc merge -> chi dong dau nhom co gia tri
        feature = str(case.get("Feature") or "").strip()
        if feature:
            current_feature = feature
        case["feature"] = current_feature
        cases.append(case)

    return cases


def resolve_pytest_targets(feature_name: str) -> list[str]:
    return list(FEATURE_TARGETS.get((feature_name or "").strip(), []))


def _run(targets: list[str], extra: list[str], dry_run: bool) -> int:
    command = [sys.executable, "-m", "pytest", *targets, *extra]
    print(f"[RUN] {' '.join(command)}")
    if dry_run:
        return 0
    return subprocess.run(command, cwd=str(PROJECT_ROOT), check=False).returncode


def run_by_feature(
    feature_name: str,
    excel_path: str | Path = DEFAULT_EXCEL,
    dry_run: bool = False,
    extra: list[str] | None = None,
) -> int:
    targets = resolve_pytest_targets(feature_name)
    if not targets:
        print(
            f"[ERROR] Feature khong co mapping sang pytest: {feature_name!r}. "
            f"Bo sung vao FEATURE_TARGETS."
        )
        return 3  # exit code != 0 de CI khong bao xanh gia
    return _run(targets, extra or [], dry_run)


def validate_mapping(excel_path: str | Path = DEFAULT_EXCEL) -> int:
    """Kiem tra moi Feature trong Excel deu co mapping. Dung cho CI."""
    cases = load_cases_from_excel(excel_path)
    features = []
    for case in cases:
        if case["feature"] and case["feature"] not in features:
            features.append(case["feature"])

    unmapped = [f for f in features if f not in FEATURE_TARGETS]
    orphan = [f for f in FEATURE_TARGETS if f not in features]

    print(f"[CHECK] {len(cases)} testcase / {len(features)} feature trong Excel")
    if unmapped:
        print(f"[CHECK] THIEU mapping cho: {unmapped}")
    if orphan:
        print(f"[CHECK] Mapping thua (khong co trong Excel): {orphan}")
    if not unmapped and not orphan:
        print("[CHECK] Mapping khop hoan toan.")
    return 1 if unmapped else 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Chay Genius web automation theo file testcase Excel."
    )
    parser.add_argument("--excel", default=str(DEFAULT_EXCEL))
    parser.add_argument("--feature", help="Chay 1 nhom feature, vi du: 'Home'")
    parser.add_argument("--tc", type=int, help="Chay 1 testcase theo so N°")
    parser.add_argument("-m", "--marker", help="Loc theo marker, vi du: 'not credit'")
    parser.add_argument("--list", action="store_true", help="Chi in danh sach TC")
    parser.add_argument("--check", action="store_true", help="Kiem tra mapping")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--write-results",
        action="store_true",
        help="Do ket qua tu reports/results.json vao cot PASS/FAIL",
    )
    args = parser.parse_args()

    excel_path = Path(args.excel)

    if args.check:
        return validate_mapping(excel_path)

    cases = load_cases_from_excel(excel_path)

    if args.list:
        for case in cases:
            print(
                f"#{case['n']:>3}  {case['feature']:<28} | "
                f"{case.get('Sub-scenario')}"
            )
        print(f"\nTong: {len(cases)} testcase")
        return 0

    extra = ["-m", args.marker] if args.marker else []

    if args.tc:
        match = next((c for c in cases if c["n"] == args.tc), None)
        if not match:
            print(f"[ERROR] Khong tim thay TC #{args.tc} trong Excel")
            return 2
        code = _run(
            resolve_pytest_targets(match["feature"]),
            [*extra, "-k", f"tc{args.tc:02d}"],
            args.dry_run,
        )
    elif args.feature:
        code = run_by_feature(args.feature, excel_path, args.dry_run, extra)
    else:
        code = _run(["tests"], extra, args.dry_run)

    if args.write_results and not args.dry_run:
        from results_writer import write_results

        write_results(excel_path)

    return code


if __name__ == "__main__":
    raise SystemExit(main())
